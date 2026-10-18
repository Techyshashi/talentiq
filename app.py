"""TalentIQ - AI Recruitment Tool (Streamlit + Azure AI Foundry)."""
import io
import json
import os
from concurrent.futures import ThreadPoolExecutor, as_completed

import pandas as pd
import streamlit as st
from azure.core.exceptions import AzureError, ClientAuthenticationError
from azure.identity import DefaultAzureCredential, get_bearer_token_provider
from docx import Document
from dotenv import load_dotenv
from openai import APIError, AzureOpenAI
from pypdf import PdfReader

load_dotenv()

PROJECT_ENDPOINT = os.getenv("AZURE_AI_PROJECT_ENDPOINT")
ENDPOINT = os.getenv("AZURE_OPENAI_ENDPOINT")
DEPLOYMENT = os.getenv("AZURE_OPENAI_DEPLOYMENT")
API_KEY = os.getenv("AZURE_OPENAI_API_KEY")
API_VERSION = os.getenv("AZURE_OPENAI_API_VERSION", "2024-10-21")
REASONING_EFFORT = os.getenv("REASONING_EFFORT")
APP_USERNAME = os.getenv("APP_USERNAME")
APP_PASSWORD = os.getenv("APP_PASSWORD")
_client = None

st.set_page_config(page_title="TalentIQ", page_icon="🎯", layout="wide")


# ---------------------------------------------------------------- AI helpers
def get_client():
    global _client
    if _client is not None:
        return _client
    if API_KEY and ENDPOINT:
        _client = AzureOpenAI(
            azure_endpoint=ENDPOINT, api_key=API_KEY, api_version=API_VERSION
        )
        return _client
    if API_KEY and not ENDPOINT:
        raise RuntimeError("AZURE_OPENAI_ENDPOINT is missing for API-key mode.")
    if PROJECT_ENDPOINT:
        try:
            from azure.ai.projects import AIProjectClient
        except ImportError as exc:
            raise RuntimeError(
                "azure-ai-projects is not installed. Install the approved requirements "
                "before using AZURE_AI_PROJECT_ENDPOINT."
            ) from exc
        project = AIProjectClient(
            endpoint=PROJECT_ENDPOINT, credential=DefaultAzureCredential()
        )
        _client = project.get_openai_client()
        return _client
    if not ENDPOINT:
        raise RuntimeError(
            "Missing Azure endpoint. Set AZURE_AI_PROJECT_ENDPOINT, or set "
            "AZURE_OPENAI_ENDPOINT together with AZURE_OPENAI_API_KEY."
        )
    provider = get_bearer_token_provider(
        DefaultAzureCredential(), "https://cognitiveservices.azure.com/.default"
    )
    _client = AzureOpenAI(
        azure_endpoint=ENDPOINT, azure_ad_token_provider=provider, api_version=API_VERSION
    )
    return _client


def ask(system: str, user: str, as_json: bool = False):
    kwargs = {"response_format": {"type": "json_object"}} if as_json else {}
    if REASONING_EFFORT:
        kwargs["reasoning_effort"] = REASONING_EFFORT
    if not DEPLOYMENT:
        raise RuntimeError("AZURE_OPENAI_DEPLOYMENT is missing.")
    for attempt in range(2):
        resp = get_client().chat.completions.create(
            model=DEPLOYMENT,
            messages=[{"role": "system", "content": system}, {"role": "user", "content": user}],
            **kwargs,
        )
        text = resp.choices[0].message.content or ""
        if not as_json:
            return text
        try:
            return json.loads(text)
        except json.JSONDecodeError as exc:
            if attempt == 1:
                raise ValueError(
                    "The model returned invalid JSON twice. Please try again."
                ) from exc
    raise RuntimeError("AI request did not return a response.")


def display_ai_error(prefix: str, error: Exception) -> None:
    if isinstance(error, ClientAuthenticationError):
        st.error(
            f"{prefix} Azure authentication is unavailable. Run "
            "`az login` in PowerShell, then restart Streamlit."
        )
        return
    st.error(f"{prefix} {error}")


def ask_bing_market_insights(jd: dict):
    """Use a preconfigured Foundry agent when all Bing settings and SDK support exist."""
    bing_connection_id = os.getenv("BING_PROJECT_CONNECTION_ID")
    agent_name = os.getenv("FOUNDRY_AGENT_NAME")
    if not PROJECT_ENDPOINT or not bing_connection_id or not agent_name:
        return None
    try:
        from azure.ai.projects import AIProjectClient
        from azure.ai.projects.models import (
            BingGroundingSearchConfiguration,
            BingGroundingSearchToolParameters,
            BingGroundingTool,
            PromptAgentDefinition,
        )
    except ImportError:
        return None

    credential = DefaultAzureCredential()
    project = AIProjectClient(endpoint=PROJECT_ENDPOINT, credential=credential)
    tool = BingGroundingTool(
        bing_grounding=BingGroundingSearchToolParameters(
            search_configurations=[
                BingGroundingSearchConfiguration(
                    project_connection_id=bing_connection_id
                )
            ]
        )
    )
    version = None
    try:
        version = project.agents.create_version(
            agent_name=agent_name,
            definition=PromptAgentDefinition(
                model=DEPLOYMENT,
                instructions=(
                    "Find current hiring-market insights for the role in India. "
                    "Return only valid JSON with an insights list. Cite the URLs returned by Bing."
                ),
                tools=[tool],
            ),
        )
        openai_client = project.get_openai_client()
        conversation = openai_client.conversations.create(
            items=[{
                "type": "message",
                "role": "user",
                "content": (
                    'Return JSON matching this schema: {"insights":[{"company":"",'
                    '"experience_range":"","compensation_range_inr":"","hiring_location":"",'
                    '"source":"","is_estimate":false}]}. '
                    f"Role details: {json.dumps(jd)}"
                ),
            }]
        )
        response = openai_client.responses.create(
            conversation=conversation.id,
            extra_body={"agent_reference": {"name": version.name, "type": "agent_reference"}},
        )
        return json.loads(response.output_text).get("insights", [])
    except (APIError, AzureError, AttributeError, RuntimeError, TypeError, ValueError):
        return None
    finally:
        if version is not None:
            project.agents.delete_version(
                agent_name=agent_name, agent_version=version.version, force=True
            )


def rank_one_cv(jd_json: str, filename: str, text: str) -> dict:
    if not text.strip():
        return {
            "Candidate (file)": filename,
            "Score /100": 0,
            "Reason": "No readable text was extracted from this CV.",
            "Matched skills": "",
            "Missing skills": "",
            "Evidence from CV": "",
        }
    try:
        result = ask(
            RANK_SYSTEM,
            f"JOB REQUIREMENTS:\n{jd_json}\n\nCV TEXT:\n{text[:12000]}",
            as_json=True,
        )
        return {
            "Candidate (file)": filename,
            "Score /100": max(0, min(100, int(result.get("score", 0)))),
            "Reason": result.get("reason", ""),
            "Matched skills": ", ".join(result.get("matched_skills", [])),
            "Missing skills": ", ".join(result.get("missing_skills", [])),
            "Evidence from CV": result.get("evidence", ""),
        }
    except (
        APIError,
        AzureError,
        AttributeError,
        KeyError,
        RuntimeError,
        TypeError,
        ValueError,
    ) as exc:
        return {
            "Candidate (file)": filename,
            "Score /100": 0,
            "Reason": f"Error: {exc}",
            "Matched skills": "",
            "Missing skills": "",
            "Evidence from CV": "",
        }


# ------------------------------------------------------------- file reading
def read_file(f) -> str:
    name = f.name.lower()
    data = f.getvalue()
    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((p.extract_text() or "") for p in reader.pages)
    if name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        return "\n".join(parts)
    if name.endswith(".txt"):
        return data.decode("utf-8", errors="ignore")
    return ""


# ------------------------------------------------------------------ prompts
JD_SYSTEM = """You are a recruitment analyst. Extract structured data from the job description.
Return JSON with exactly these keys:
role (string), must_have_skills (list of strings), nice_to_have_skills (list of strings),
experience (string), location (string), education (string), domain (string).
If something is not stated write "Not stated". Never guess."""

RANK_SYSTEM = """You are a fair, evidence-based CV screener.
Score ONE CV against the job requirements, out of 100, using this rubric:
Skills match 40, Relevant experience 30, Domain/industry 15, Education/certifications 10, Achievements 5.
FAIRNESS RULES: use ONLY evidence written in the CV. Ignore name, gender, age, photo,
marital status, religion, nationality and address. Do not infer anything not written.
Return JSON with keys:
score (integer 0-100), reason (max 2 short sentences),
matched_skills (list), missing_skills (list), evidence (one short quote from the CV)."""

QUESTIONS_SYSTEM = """You are a senior recruiter. Create exactly 10 phone-screening questions for this job:
4 with type "Technical", 3 with type "Experience", 3 with type "Fit".
Return JSON: {"questions":[{"type":"","question":"","good_answer":"what a good answer looks like"}]}"""

MARKET_SYSTEM = """You are a hiring-market analyst for India.
List 6 to 8 companies likely hiring for a role similar to the one given.
Return JSON: {"insights":[{"company":"","experience_range":"","compensation_range_inr":"",
"hiring_location":"","source":"","is_estimate":true}]}
RULES: If the recruiter supplied public job-post text, extract ONLY from it and put the post URL or
title in "source" with is_estimate=false. If no job-post text was supplied, use general knowledge,
set is_estimate=true and set source to "Estimate - model knowledge, not verified".
Never invent URLs or fake sources."""

SEARCH_SYSTEM = """You are a sourcing expert for Indian recruitment.
Return JSON with keys: naukri_boolean, naukri_plain, linkedin_boolean, linkedin_plain.
Boolean versions must use AND, OR, NOT, quotation marks and parentheses.
Plain versions are simple keyword phrases a recruiter can type without operators."""


# ------------------------------------------------------------------- login
if "logged_in" not in st.session_state:
    st.session_state.logged_in = False

if not st.session_state.logged_in:
    st.title("🎯 TalentIQ - AI Recruitment Tool")
    st.subheader("Recruiter sign-in")
    missing_login = [
        name for name, value in (("APP_USERNAME", APP_USERNAME), ("APP_PASSWORD", APP_PASSWORD))
        if not value
    ]
    if missing_login:
        st.error("Missing .env value(s): " + ", ".join(missing_login))
        st.stop()
    with st.container(border=True):
        user = st.text_input("Username")
        pwd = st.text_input("Password", type="password")
        if st.button("Sign in", type="primary"):
            if user == APP_USERNAME and pwd == APP_PASSWORD:
                st.session_state.logged_in = True
                st.rerun()
            else:
                st.error("Invalid credentials")
    st.caption("Demo login. Credentials are set in the .env file.")
    st.stop()

# --------------------------------------------------------------- main app
st.title("🎯 TalentIQ - AI Recruitment Tool")
with st.sidebar:
    st.success("Signed in as recruiter")
    st.caption(f"Model deployment: {DEPLOYMENT}")
    if st.button("Sign out"):
        st.session_state.clear()
        st.rerun()

tabs = st.tabs(
    ["2. Job Description", "3. Upload CVs", "4. Stack Ranking",
     "5. Screening Questions", "6. Market Insights", "7. Search Strings"]
)

# ---- 2. JD
with tabs[0]:
    st.header("Upload the Job Description")
    jd_file = st.file_uploader("PDF / Word / TXT", type=["pdf", "docx", "txt"], key="jd_file")
    jd_pasted = st.text_area("...or paste the JD text", height=180)
    if st.button("Analyse JD", type="primary"):
        text = read_file(jd_file) if jd_file else jd_pasted
        if not text.strip():
            st.warning("Please upload or paste a JD first.")
        else:
            with st.spinner("Reading the JD..."):
                try:
                    st.session_state.jd_text = text
                    st.session_state.jd = ask(JD_SYSTEM, text, as_json=True)
                except Exception as e:
                    display_ai_error("AI call failed.", e)

    if "jd" in st.session_state:
        jd = st.session_state.jd
        st.subheader("What the AI understood")
        c1, c2 = st.columns(2)
        jd["role"] = c1.text_input("Role", jd.get("role", ""))
        jd["experience"] = c2.text_input("Experience", jd.get("experience", ""))
        jd["location"] = c1.text_input("Location", jd.get("location", ""))
        jd["domain"] = c2.text_input("Domain", jd.get("domain", ""))
        jd["education"] = st.text_input("Education", jd.get("education", ""))
        jd["must_have_skills"] = [s.strip() for s in st.text_area(
            "Must-have skills (comma separated - edit if the AI got it wrong)",
            ", ".join(jd.get("must_have_skills", []))).split(",") if s.strip()]
        jd["nice_to_have_skills"] = [s.strip() for s in st.text_area(
            "Nice-to-have skills", ", ".join(jd.get("nice_to_have_skills", []))).split(",") if s.strip()]
        st.session_state.jd = jd
        st.info("Edits above are used for all later steps.")

# ---- 3. CVs
with tabs[1]:
    st.header("Upload CVs")
    cv_files = st.file_uploader("Select many PDF / Word CVs at once",
                                type=["pdf", "docx"], accept_multiple_files=True)
    if cv_files and st.button("Read CVs", type="primary"):
        cvs, failed = {}, []
        for f in cv_files:
            text = read_file(f)
            (cvs.__setitem__(f.name, text) if text.strip() else failed.append(f.name))
        st.session_state.cvs = cvs
        st.session_state.failed = failed
    if "cvs" in st.session_state:
        st.success(f"Read {len(st.session_state.cvs)} of "
                   f"{len(st.session_state.cvs) + len(st.session_state.failed)} CVs")
        if st.session_state.failed:
            st.warning("Could not read (scanned image or empty?): " + ", ".join(st.session_state.failed))
        st.dataframe(pd.DataFrame({"File": list(st.session_state.cvs.keys()),
                                   "Characters read": [len(t) for t in st.session_state.cvs.values()]}),
                     use_container_width=True)

# ---- 4. Ranking
with tabs[2]:
    st.header("Stack Ranking")
    if "jd" not in st.session_state or "cvs" not in st.session_state:
        st.info("Complete steps 2 and 3 first.")
    elif st.button("Rank all CVs", type="primary"):
        rows = []
        bar = st.progress(0.0, text="Starting CV ranking...")
        jd_json = json.dumps(st.session_state.jd)
        items = list(st.session_state.cvs.items())
        if not items:
            st.warning("No readable CVs are available to rank.")
            st.stop()
        with ThreadPoolExecutor(max_workers=4) as executor:
            futures = {
                executor.submit(rank_one_cv, jd_json, fname, text): fname
                for fname, text in items
            }
            for i, future in enumerate(as_completed(futures), 1):
                result = future.result()
                if not result["Evidence from CV"] and not st.session_state.cvs[
                    futures[future]
                ].strip():
                    st.warning(f"{futures[future]} has no readable text and was not ranked.")
                rows.append(result)
                bar.progress(
                    i / len(items),
                    text=f"Completed {i} of {len(items)} CV rankings",
                )
        df = pd.DataFrame(rows).sort_values("Score /100", ascending=False).reset_index(drop=True)
        df.index += 1
        df.index.name = "Rank"
        st.session_state.ranking = df
        st.success(f"Ranking complete: {len(rows)} CVs processed.")
    if "ranking" in st.session_state:
        st.dataframe(st.session_state.ranking, use_container_width=True)
        st.download_button(
            "Download ranking as CSV",
            data=st.session_state.ranking.to_csv(index=True).encode("utf-8"),
            file_name="talentiq-ranking.csv",
            mime="text/csv",
        )
        st.caption("Scores use only CV evidence. Name, gender, age, photo and marital status are ignored. "
                   "Always verify a few scores against the CVs.")

# ---- 5. Questions
with tabs[3]:
    st.header("Screening Questions")
    if "jd" not in st.session_state:
        st.info("Complete step 2 first.")
    elif st.button("Generate 10 questions", type="primary"):
        with st.spinner("Writing questions..."):
            try:
                question_prompt = json.dumps(st.session_state.jd)
                questions = []
                for attempt in range(2):
                    questions = ask(
                        QUESTIONS_SYSTEM,
                        question_prompt,
                        as_json=True,
                    ).get("questions", [])
                    counts = {
                        question_type: sum(
                            q.get("type") == question_type for q in questions
                        )
                        for question_type in ("Technical", "Experience", "Fit")
                    }
                    if len(questions) == 10 and counts == {
                        "Technical": 4,
                        "Experience": 3,
                        "Fit": 3,
                    }:
                        break
                    question_prompt += (
                        "\nYour previous response had the wrong count. Return exactly "
                        "10 questions: 4 Technical, 3 Experience, and 3 Fit."
                    )
                if len(questions) != 10:
                    raise ValueError("The model did not return exactly 10 questions.")
                st.session_state.questions = questions
            except Exception as e:
                display_ai_error("Could not generate questions.", e)
    for i, q in enumerate(st.session_state.get("questions", []), 1):
        with st.expander(f"Q{i} [{q.get('type')}] {q.get('question')}"):
            st.markdown(f"**Good answer:** {q.get('good_answer')}")

# ---- 6. Market insights
with tabs[4]:
    st.header("Hiring Market Insights")
    st.caption("Paste text from PUBLIC job posts (company career pages, etc.) for sourced results. "
               "Leave empty to get clearly labelled estimates. Do not scrape or log in to LinkedIn/Naukri.")
    if PROJECT_ENDPOINT and os.getenv("BING_PROJECT_CONNECTION_ID") and os.getenv("FOUNDRY_AGENT_NAME"):
        st.info("Bing-grounded mode is configured through the Foundry agent. "
                "If it is unavailable, the app uses estimates or pasted public posts.")
    posts = st.text_area("Public job-post text / URLs (optional)", height=150)
    if "jd" not in st.session_state:
        st.info("Complete step 2 first.")
    elif st.button("Get market insights", type="primary"):
        with st.spinner("Analysing..."):
            try:
                data = None if posts else ask_bing_market_insights(st.session_state.jd)
                if data is None:
                    prompt = (f"ROLE DETAILS:\n{json.dumps(st.session_state.jd)}\n\n"
                              f"PUBLIC JOB POSTS SUPPLIED:\n{posts or 'NONE'}")
                    data = ask(MARKET_SYSTEM, prompt, as_json=True).get("insights", [])
                st.session_state.market = pd.DataFrame(data)
            except Exception as e:
                display_ai_error("Could not load market insights.", e)
    if "market" in st.session_state:
        st.dataframe(st.session_state.market, use_container_width=True)
        st.warning("Verify company names and salary figures before relying on them.")

# ---- 7. Search strings
with tabs[5]:
    st.header("Search Strings")
    if "jd" not in st.session_state:
        st.info("Complete step 2 first.")
    elif st.button("Generate search strings", type="primary"):
        with st.spinner("Building strings..."):
            try:
                st.session_state.strings = ask(
                    SEARCH_SYSTEM, json.dumps(st.session_state.jd), as_json=True
                )
            except Exception as e:
                display_ai_error("Could not generate search strings.", e)
    s = st.session_state.get("strings")
    if s:
        st.subheader("Naukri")
        st.markdown("**Boolean**"); st.code(s.get("naukri_boolean", ""), language=None)
        st.markdown("**Plain**"); st.code(s.get("naukri_plain", ""), language=None)
        st.subheader("LinkedIn")
        st.markdown("**Boolean**"); st.code(s.get("linkedin_boolean", ""), language=None)
        st.markdown("**Plain**"); st.code(s.get("linkedin_plain", ""), language=None)