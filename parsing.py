"""Safe text extraction and upload validation for TalentIQ."""
from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from pypdf import PdfReader

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}
MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_TEXT_CHARS = 50_000


def validate_upload(filename: str, data: bytes, cv: bool = False) -> None:
    ext = Path(filename or "").suffix.lower()
    allowed = {".pdf", ".docx"} if cv else ALLOWED_EXTENSIONS
    if ext not in allowed:
        raise ValueError(f"Unsupported file type. Allowed: {', '.join(sorted(allowed))}")
    if not data:
        raise ValueError("The uploaded file is empty.")
    if len(data) > MAX_FILE_BYTES:
        raise ValueError("The uploaded file exceeds the 10 MB limit.")


def extract_text(filename: str, data: bytes) -> str:
    validate_upload(filename, data)
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)
    elif ext == ".docx":
        document = Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs]
        parts.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        text = "\n".join(parts)
    else:
        text = data.decode("utf-8", errors="ignore")
    return text[:MAX_TEXT_CHARS].strip()
